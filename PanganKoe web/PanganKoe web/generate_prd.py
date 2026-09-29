import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn
import os

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_table_borders(table, color="CCCCCC"):
    tblPr = table._element.xpath('w:tblPr')
    if tblPr:
        borders = parse_xml(f'<w:tblBorders {nsdecls("w")}><w:top w:val="single" w:sz="4" w:space="0" w:color="{color}"/><w:bottom w:val="single" w:sz="4" w:space="0" w:color="{color}"/><w:insideH w:val="single" w:sz="4" w:space="0" w:color="{color}"/><w:insideV w:val="none"/><w:left w:val="none"/><w:right w:val="none"/></w:tblBorders>')
        tblPr[0].append(borders)

def create_prd():
    doc = docx.Document()
    
    # Page Margins
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        
    # Styles
    style_normal = doc.styles['Normal']
    style_normal.font.name = 'Arial'
    style_normal.font.size = Pt(10.5)
    style_normal.font.color.rgb = RGBColor(30, 41, 59) # Slate 800

    # Colors
    C_PRIMARY = RGBColor(5, 150, 105) # Emerald 600
    C_DARK = RGBColor(15, 23, 42)     # Slate 900
    C_MUTED = RGBColor(100, 116, 139)  # Slate 500

    # Title Banner / Header Cover
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = p_title.add_run("PRODUCT REQUIREMENT DOCUMENT (PRD)\n")
    run_title.bold = True
    run_title.font.size = Pt(22)
    run_title.font.color.rgb = C_PRIMARY

    run_sub = p_title.add_run("PANGANKOE - Meal Planner & Connector to UMKM\n")
    run_sub.bold = True
    run_sub.font.size = Pt(14)
    run_sub.font.color.rgb = C_DARK

    run_desc = p_title.add_run("Spesifikasi Sistem & Dokumen Kebutuhan Produk (Web Application)\nFinal Project PFB 2026 - Kelompok 5")
    run_desc.font.size = Pt(11)
    run_desc.font.color.rgb = C_MUTED
    doc.add_paragraph() # spacing

    # Table Metadata Document
    table_meta = doc.add_table(rows=6, cols=2)
    table_meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table_meta)
    
    meta_data = [
      ("Nama Project", "PanganKoe (Meal Planner & Connector to UMKM)"),
      ("Versi Dokumen", "v1.0.0 (Final Approved for Production Development)"),
      ("Tanggal Rilis PRD", "September 2026"),
      ("Tim Penulis / Pengembang", "Kelompok 5 (Alfino, Fadil, Khalifa, Keane, Matthew)"),
      ("Platform Target", "Desktop Web, Tablet & Mobile Web Responsive"),
      ("Status Integrasi Database", "Siap Terhubung dengan phpMyAdmin / MySQL (`pangankoe_db`)")
    ]
    
    for idx, (k, v) in enumerate(meta_data):
        row = table_meta.rows[idx]
        cell_k, cell_v = row.cells[0], row.cells[1]
        
        cell_k.width = Inches(2.2)
        cell_v.width = Inches(4.3)
        
        pk = cell_k.paragraphs[0]
        rk = pk.add_run(k)
        rk.bold = True
        rk.font.color.rgb = C_PRIMARY
        
        pv = cell_v.paragraphs[0]
        pv.add_run(v)
        
        set_cell_background(cell_k, "ECFDF5")
        set_cell_background(cell_v, "F8FAFC")
        set_cell_margins(cell_k, 80, 80, 100, 100)
        set_cell_margins(cell_v, 80, 80, 100, 100)

    doc.add_paragraph() # Spacing

    # Helper function for Headings
    def add_heading_1(text):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(18)
        h.paragraph_format.space_after = Pt(6)
        r = h.add_run(text)
        r.bold = True
        r.font.size = Pt(14)
        r.font.color.rgb = C_PRIMARY
        return h

    def add_heading_2(text):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(12)
        h.paragraph_format.space_after = Pt(4)
        r = h.add_run(text)
        r.bold = True
        r.font.size = Pt(12)
        r.font.color.rgb = C_DARK
        return h

    def add_bullet(text, bold_prefix=""):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(3)
        if bold_prefix:
            r_b = p.add_run(bold_prefix + " ")
            r_b.bold = True
            r_b.font.color.rgb = C_DARK
        p.add_run(text)

    # 1. EXECUTIVE SUMMARY & VISION
    add_heading_1("1. RINGKASAN EKSEKUTIF & VISI PRODUK")
    p = doc.add_paragraph()
    r_bold = p.add_run("PanganKoe ")
    r_bold.bold = True
    p.add_run("adalah platform aplikasi berbasis web dan mobile yang hadir sebagai solusi terintegrasi atas tiga permasalahan utama dalam masyarakat Indonesia: kebingungan perencanaan menu harian rumah tangga, krisis gizi & stunting nasional, serta keterbatasan akses pasar digital bagi UMKM kuliner lokal.")
    
    add_heading_2("Visi Utama Produk:")
    add_bullet("Menjadi sistem perencanaan pangan keluarga paling cerdas, personal, dan terjangkau di Indonesia yang terhubung langsung dengan ekosistem ekonomi warung lokal sekitar.", "Visi:")
    add_bullet("Menyediakan rekomendasi menu makan bergizi seimbang (AKG) dalam waktu kurang dari 30 detik yang disesuaikan dengan anggaran dapur keluarga.", "Misi 1:")
    add_bullet("Mengotomatiskan kalkulasi daftar belanja harian/mingguan untuk mengeliminasi pemborosan makanan (anti food-waste).", "Misi 2:")
    add_bullet("Memfasilitasi digitalisasi UMKM kuliner lokal melalui integrasi transaksi mudah berbasis WhatsApp.", "Misi 3:")

    # 2. PROBLEM STATEMENT & MARKET OPPORTUNITY
    add_heading_1("2. PERNYATAAN MASALAH & PELUANG PASAR")
    p = doc.add_paragraph()
    p.add_run("Berdasarkan data empiris Survei Status Gizi Indonesia (SSGI) 2024 serta data Kadin & Kemenkop RI 2024, ditemukan kesenjangan besar dalam konsumsi gizi dan digitalisasi UMKM:")
    
    add_bullet("Prevalensi stunting nasional masih mencapai 19,8% (hampir 1 dari 5 balita di Indonesia mengalami gangguan tumbuh kembang).", "Masalah Gizi & Stunting:")
    add_bullet("Keragaman pangan rata-rata keluarga Indonesia baru mencapai 44,6%, jauh di bawah standar ideal kecukupan gizi harian.", "Keragaman Pangan Rendah:")
    add_bullet("Ibu rumah tangga dan anak kos sering mengalami kendala kebingungan menu harian dan ketidakpastian anggaran dapur.", "Pemborosan Belanja:")
    add_bullet("Terdapat lebih dari 6,4 juta unit UMKM kuliner di Indonesia, namun >60% belum terdigitalisasi karena komisi aplikasi tinggi dan kerumitan teknologi.", "Kesenjangan Digital UMKM:")

    # 3. SDGS ALIGNMENT
    add_heading_1("3. KESELARASAN DENGAN SDGS (4 PILAR KEBERLANJUTAN)")
    p = doc.add_paragraph("Aplikasi PanganKoe secara langsung berkontribusi pada 4 Sustainable Development Goals (SDGs):")
    
    table_sdg = doc.add_table(rows=5, cols=3)
    table_sdg.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table_sdg)
    
    headers = ["SDG Target", "Nama Tujuan", "Kontribusi Konkret PanganKoe"]
    hdr_cells = table_sdg.rows[0].cells
    for i, h_text in enumerate(headers):
        r_hdr = hdr_cells[i].paragraphs[0].add_run(h_text)
        r_hdr.bold = True
        r_hdr.font.color.rgb = RGBColor(255, 255, 255)
        set_cell_background(hdr_cells[i], "059669")
        set_cell_margins(hdr_cells[i], 100, 100, 120, 120)

    sdg_data = [
      ("SDG 2", "Tanpa Kelaparan", "Mendorong keragaman pangan harian & berkontribusi pada penurunan stunting nasional via Meal Planner gizi mikro."),
      ("SDG 3", "Kehidupan Sehat & Sejahtera", "Pemantauan kecukupan gizi harian (AKG) untuk mencegah penyakit tidak menular dan malnutrisi."),
      ("SDG 8", "Pekerjaan Layak & Pertumbuhan Ekonomi", "Memberdayakan UMKM kuliner & pedagang sayur lokal dengan akses pasar digital ramah WhatsApp."),
      ("SDG 9", "Inovasi & Infrastruktur", "Menghadirkan inovasi digitalisasi sektor pangan yang mudah diakses seluruh lapisan masyarakat.")
    ]

    for idx, (s, n, k) in enumerate(sdg_data, start=1):
        row_cells = table_sdg.rows[idx].cells
        r_s = row_cells[0].paragraphs[0].add_run(s)
        r_s.bold = True
        r_n = row_cells[1].paragraphs[0].add_run(n)
        r_n.bold = True
        row_cells[2].paragraphs[0].add_run(k)
        bg = "F8FAFC" if idx % 2 == 1 else "ECFDF5"
        for c in row_cells:
            set_cell_background(c, bg)
            set_cell_margins(c, 80, 80, 100, 100)

    doc.add_paragraph()

    # 4. TARGET USER PERSONAS
    add_heading_1("4. TARGET USER PERSONAS & SPESIFIKASI SEGMEN")
    
    add_heading_2("Segmen Konsumen (Pengguna Utama):")
    add_bullet("Ibu rumah tangga usia 25–45 tahun pengelola menu dan keuangan dapur keluarga (Prioritas Utama untuk nutrisi anak balita & pencegahan stunting).", "Pengguna Primer (IRT):")
    add_bullet("Anak kos, mahasiswa, dan pekerja muda yang membutuhkan makanan bergizi dengan budget hemat dan praktis.", "Pengguna Sekunder (Anak Kos/Muda):")
    add_bullet("Orang tua balita dan individu dengan kebutuhan gizi khusus (pencegahan stunting & gizi medis).", "Pengguna Tersier (Parental Care):")

    add_heading_2("Segmen Mitra UMKM:")
    add_bullet("Penjual masakan siap saji rumahan skala RT/RW hingga kelurahan.", "Warung Makan Lokal:")
    add_bullet("Toko sayur keliling, pedagang protein (ayam, ikan, tempe), dan toko sembako.", "Toko Bahan Makanan:")
    add_bullet("Ibu-ibu penyedia katering rumahan harian.", "Katering Rumahan:")

    # 5. PRODUCT SCOPE & FUNCTIONAL REQUIREMENTS
    add_heading_1("5. CAKUPAN PRODUK & PERSYARATAN FUNGSIONAL (FR)")
    p = doc.add_paragraph("Berikut adalah rincian kebutuhan fungsional (Functional Requirements) yang telah diimplementasikan dalam prototype web PanganKoe:")

    fr_list = [
      ("FR-01: Web Splash Screen Overlay", "Menampilkan animasi logo pulse, progress loader (0-100%), tagline 'Meal Planner & Connector UMKM', serta tombol skip untuk transisi cepat ke website."),
      ("FR-02: Web Onboarding Showcase", "Interaktif 3-Step Feature Tabs (Smart Meal Planner, Kalkulator Belanja Efisien, Konektor UMKM) lengkap dengan hero metrics dan preview card."),
      ("FR-03: Dual Auth Modal (Sign In & Sign Up)", "Modal autentikasi desktop split-screen. Mendukung input Email/WA, Kata Sandi, Role Selection (IRT, Anak Kos, UMKM), Password Strength Meter, Checkbox Syarat & UU PDP 2022, serta Social Login (Google & WA OTP)."),
      ("FR-04: Redirection ke Halaman Beranda", "Setelah autentikasi sukses (login/register), sistem secara otomatis mengarahkan pengguna ke Halaman Beranda Dashboard dengan memuat Nama & Peran Pengguna secara dinamis."),
      ("FR-05: Smart Meal Planner Module (<30 Dt)", "Rekomendasi menu Sarapan, Makan Siang (Menu Utama Stunting), dan Makan Malam. Menghitung otomatis Kalori, Protein, Zat Besi (Fe), Fiber, dan Omega-3 disesuaikan standar AKG."),
      ("FR-06: Automated Grocery Calculator", "Mengelompokkan bahan makanan per kategori (Protein, Sayur, Nabati, Bumbu). Dilengkapi checklist interaktif yang menghitung sisa anggaran belanja secara live."),
      ("FR-07: Direktori Mitra UMKM & WA Order", "Menampilkan daftar mitra warung/toko terdekat (radius 150m-500m) dilengkapi tombol pesanan langsung berbasis link WhatsApp."),
      ("FR-08: Monitoring Dampak SDGs", "Dashboard visualisasi kontribusi pengguna terhadap pencapaian SDG 2, SDG 3, SDG 8, dan SDG 9.")
    ]

    for code_title, desc in fr_list:
        add_bullet(desc, code_title + ":")

    # 6. TECHNICAL ARCHITECTURE & DATABASE SPECIFICATIONS
    add_heading_1("6. ARSITEKTUR TEKNIS & SPESIFIKASI DATABASE (phpMyAdmin / MySQL)")
    p = doc.add_paragraph("Aplikasi didesain menggunakan arsitektur Web modern yang siap dihubungkan langsung ke database phpMyAdmin/MySQL (`pangankoe_db`):")

    add_heading_2("Spesifikasi Skema Tabel `users` (MySQL/phpMyAdmin):")

    table_db = doc.add_table(rows=7, cols=4)
    table_db.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table_db)

    db_headers = ["Field Name", "Data Type", "Constraints", "Keterangan"]
    for i, h_text in enumerate(db_headers):
        cell = table_db.rows[0].cells[i]
        r_dbh = cell.paragraphs[0].add_run(h_text)
        r_dbh.bold = True
        r_dbh.font.color.rgb = RGBColor(255, 255, 255)
        set_cell_background(cell, "059669")
        set_cell_margins(cell, 80, 80, 100, 100)

    db_schema = [
      ("id", "INT(11)", "PRIMARY KEY, AUTO_INCREMENT", "ID Unik Pengguna"),
      ("full_name", "VARCHAR(100)", "NOT NULL", "Nama Lengkap Pengguna"),
      ("email", "VARCHAR(100)", "NOT NULL, UNIQUE", "Email / No. WhatsApp (Identity)"),
      ("user_role", "ENUM", "NOT NULL, DEFAULT 'IRT'", "'IRT', 'AnakKos', 'UMKM', 'Lainnya'"),
      ("password", "VARCHAR(255)", "NOT NULL", "Hash Kata Sandi (BCRYPT)"),
      ("created_at", "TIMESTAMP", "DEFAULT CURRENT_TIMESTAMP", "Waktu Pendaftaran Akun")
    ]

    for idx, (f, t, c, k) in enumerate(db_schema, start=1):
        r_cells = table_db.rows[idx].cells
        r_f = r_cells[0].paragraphs[0].add_run(f)
        r_f.bold = True
        r_cells[1].paragraphs[0].add_run(t)
        r_cells[2].paragraphs[0].add_run(c)
        r_cells[3].paragraphs[0].add_run(k)
        bg = "F8FAFC" if idx % 2 == 1 else "ECFDF5"
        for cell in r_cells:
            set_cell_background(cell, bg)
            set_cell_margins(cell, 60, 60, 80, 80)

    doc.add_paragraph()

    add_heading_2("Integrasi Endpoints API (PHP):")
    add_bullet("Menerima input identity & password, memverifikasi hash password di database, dan mengembalikan status JSON success.", "POST /login.php:")
    add_bullet("Menerima full_name, email, user_role, & password, melakukan enkripsi hash, dan menyimpan data user baru ke database MySQL.", "POST /register.php:")
    add_bullet("Membuat koneksi MySQLi / PDO ke database phpMyAdmin 'pangankoe_db'.", "connection.php:")

    # 7. NON-FUNCTIONAL REQUIREMENTS
    add_heading_1("7. PERSYARATAN NON-FUNGSIONAL (NFR)")
    add_bullet("Waktu muat halaman (Page Load) < 1.5 detik; Waktu pembuatan Meal Plan < 30 detik.", "Performa (Performance):")
    add_bullet("Enkripsi kata sandi menggunakan BCRYPT; Kepatuhan terhadap UU Republik Indonesia No. 27 Tahun 2022 tentang Perlindungan Data Pribadi (PDP).", "Keamanan (Security):")
    add_bullet("Desain antarmuka intuitif dengan sistem warna kontras tinggi (WCAG 2.1 AA) dan kemudahan navigasi di layar smartphone maupun PC.", "Usability & Accessibility:")
    add_bullet("Mampu menangani hingga 10.000 pengguna aktif harian dengan arsitektur database terindeks.", "Scalability:")

    # 8. UI/UX DESIGN SYSTEM SPECIFICATIONS
    add_heading_1("8. SPESIFIKASI DESAIN SYSTEM & SPESIFIKASI WARNA COLOR PALETTE")
    p = doc.add_paragraph("Antarmuka aplikasi menggunakan sistem warna minimalis bernuansa Light Green & Mint:")

    add_bullet("#10B981 (Emerald 500) & #059669 (Emerald 600) — Memberikan kesan segar, sehat, dan alami.", "Primary Color:")
    add_bullet("#ECFDF5 (Emerald 50) & #D1FAE5 (Emerald 100) — Latar belakang lembut yang nyaman di mata.", "Background & Mint Accent:")
    add_bullet("#FFFFFF dengan efek Glassmorphism (Backdrop-filter blur 12px) & Border #D1FAE5.", "Card & Surface:")
    add_bullet("#0F172A (Slate 900) & #1E293B (Slate 800) — Memastikan keterbacaan teks maksimal.", "Typography Color:")
    add_bullet("Google Font Plus Jakarta Sans (Weight 400, 500, 600, 700, 800).", "Font Family:")

    # 9. SUCCESS METRICS & KPIS
    add_heading_1("9. METRIK KEBERHASILAN PRODUK (KPIS)")
    add_bullet(">85% pengguna berhasil menyusun meal plan mingguan dalam waktu < 30 detik pada sesi pertama.", "Adopsi Meal Planner:")
    add_bullet("Penurunan indeks pemborosan bahan makanan rumah tangga pengguna hingga 35-40%.", "Efisiensi Belanja:")
    add_bullet("Pendaftaran 500+ mitra UMKM kuliner lokal dalam 3 bulan pertama rilis.", "Pemberdayaan UMKM:")

    # 10. IMPLEMENTATION ROADMAP
    add_heading_1("10. ROADMAP IMPLEMENTASI & FASE RELEASE")
    add_bullet("Desain & Pengembangan Prototype Frontend (Splash, Onboarding, Auth Modal, Halaman Beranda Dashboard) — SELESAI (100%).", "Fase 1 (Frontend Development):")
    add_bullet("Pengembangan Database phpMyAdmin MySQL & Script PHP API (login.php, register.php, connection.php) — SIAP DIHUBUNGKAN.", "Fase 2 (Backend & DB Integration):")
    add_bullet("Integrasi API WhatsApp Gateway & Dashboard Mitra UMKM.", "Fase 3 (Merchant Integration):")
    add_bullet("Pengujian QA, Security Audit UU PDP, & Rilis Produksi.", "Fase 4 (QA & Launch):")

    doc.add_paragraph()
    
    # Save files to paths
    path1 = r"D:\Semester 3\Programming for business\Etc folder\FINAL PROJECT PANGANKOE\PRD_PanganKoe_Web_Application.docx"
    path2 = r"D:\Semester 3\Programming for business\Etc folder\FINAL PROJECT PANGANKOE\PanganKoe web\PRD_PanganKoe_Web_Application.docx"
    
    doc.save(path1)
    doc.save(path2)
    print(f"PRD successfully created at:\n1. {path1}\n2. {path2}")

if __name__ == "__main__":
    create_prd()
